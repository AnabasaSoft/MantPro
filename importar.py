# -*- coding: utf-8 -*-
# SPDX-License-Identifier: AGPL-3.0-or-later
#
# MantPro - Sistema de Mantenimiento Preventivo
# Copyright (C) 2026 AnabasaSoft <anabasasoft@gmail.com>
#
# Este programa es software libre: puedes redistribuirlo y/o modificarlo bajo
# los términos de la Licencia Pública General Affero de GNU publicada por la
# Free Software Foundation, ya sea la versión 3 de la Licencia o (a tu
# elección) cualquier versión posterior.
#
# Este programa se distribuye con la esperanza de que sea útil, pero SIN
# NINGUNA GARANTÍA; ni siquiera la garantía implícita de COMERCIABILIDAD o
# IDONEIDAD PARA UN PROPÓSITO PARTICULAR. Consulta la Licencia Pública
# General Affero de GNU para más detalles.
#
# Deberías haber recibido una copia de la Licencia Pública General Affero de
# GNU junto con este programa. Si no, consulta <https://www.gnu.org/licenses/>.

"""Importación de datos desde Excel (.xlsx) o CSV: máquinas, artículos del
almacén, avisos recurrentes y usuarios.

No depende de Qt (la ventana del asistente está en main.py), para que encaje
con la futura separación motor/interfaz del modo servidor.

La vista previa y la importación real usan exactamente el mismo código
(procesar): se hace todo dentro de una transacción y, en la vista previa, se
deshace al final con ROLLBACK. Así lo que se enseña en la vista previa es lo
que va a pasar, incluidas las zonas, estanterías o baldas que se crearán.
"""

import csv
import re
import secrets
import sqlite3
import unicodedata
from datetime import date, datetime, timedelta

import almacen
import idiomas
import usuarios
from idiomas import t

TIPOS = ("maquinas", "articulos", "avisos", "usuarios")

# (campo, obligatorio) de cada tipo, en el orden en que salen en el asistente
# y en la plantilla. La etiqueta visible de cada campo es t("imp_campo_<campo>").
CAMPOS = {
    "maquinas": [("nombre", True), ("padre", False), ("zona", False), ("notas", False)],
    "articulos": [("codigo", False), ("nombre", True), ("descripcion", False), ("unidad", False),
                  ("stock", False), ("minimo", False), ("zona", False), ("estanteria", False),
                  ("balda", False), ("seccion", False)],
    "avisos": [("titulo", True), ("fecha_inicio", True), ("frecuencia", True), ("margen", False),
               ("tecnico", False), ("especialidad", False), ("ultima", False)],
    "usuarios": [("login", True), ("nombre", True), ("roles", False), ("especialidad", False)],
}

MARGEN_POR_DEFECTO = 5  # igual que el alta de avisos en la pestaña Avisos
MOTIVO_AJUSTE_IMPORTACION = almacen.MOTIVO_AJUSTE_IMPORTACION

# Sinónimos (además de la etiqueta del campo en los 7 idiomas) para emparejar
# solas las columnas de un Excel que no es nuestra plantilla. Se comparan ya
# normalizados (minúsculas, sin tildes ni signos, sin lo que va entre paréntesis).
_SINONIMOS = {
    ("maquinas", "nombre"): ["maquina", "equipo", "activo", "machine", "equipment", "asset", "makina",
                             "makinaren izena", "maquina nom", "nom de la maquina", "nombre de la maquina",
                             "nome da maquina", "nom de la machine"],
    ("maquinas", "padre"): ["padre", "parent", "maquina padre", "linea", "line", "depende de"],
    ("articulos", "nombre"): ["material", "articulo", "item", "producto", "product", "article", "artigo",
                              "artikulua", "denominacion", "description corta"],
    ("articulos", "codigo"): ["code", "ref", "referencia", "reference", "sku", "kodea", "codi", "codigo"],
    ("articulos", "stock"): ["cantidad", "existencias", "quantity", "qty", "kopurua", "quantitat",
                             "cantidade", "quantidade", "quantite", "stock actual", "current stock"],
    ("articulos", "minimo"): ["minimo", "min", "minimum", "stock min", "stock minimo"],
    ("articulos", "unidad"): ["ud", "uds", "unit", "units", "unitat", "unidade", "unite"],
    ("articulos", "estanteria"): ["estante", "shelving", "rack", "shelf unit", "apalategia", "prestatgeria",
                                  "andel", "rayonnage"],
    ("articulos", "balda"): ["shelf", "nivel", "level", "apala", "prestatge", "prateleira", "tablette"],
    ("avisos", "titulo"): ["aviso", "tarea", "task", "title", "preventivo", "abisua", "avis", "rappel", "lembrete"],
    ("avisos", "fecha_inicio"): ["inicio", "fecha", "start", "start date", "date", "data"],
    ("avisos", "frecuencia"): ["frequency", "periodicidad", "periodicity", "maiztasuna", "freqüencia",
                               "frecuencia", "frequencia", "frequence"],
    ("avisos", "margen"): ["margen", "dias", "duracion", "days", "margin", "duration"],
    ("avisos", "tecnico"): ["tecnico", "asignado", "asignado a", "technician", "assigned to", "responsable"],
    ("avisos", "ultima"): ["ultima", "ultima vez", "last done", "last", "hecho"],
    ("usuarios", "login"): ["login", "user", "username", "usuario", "erabiltzailea", "usuari", "usuario"],
    ("usuarios", "nombre"): ["nombre", "name", "nombre completo", "full name", "izena", "nom", "nome"],
    ("usuarios", "roles"): ["rol", "role", "roles", "perfil", "profile"],
}

_CLAVES_FRECUENCIA = ("freq_diario", "freq_semanal", "freq_mensual", "freq_trimestral",
                      "freq_semestral", "freq_anual")
_SINONIMOS_FRECUENCIA = {
    "freq_diario": ["diaria", "daily", "day", "egunero", "diari", "quotidien"],
    "freq_semanal": ["weekly", "week", "astero", "setmanal", "hebdomadaire"],
    "freq_mensual": ["monthly", "month", "hilero", "mensuel"],
    "freq_trimestral": ["quarterly", "quarter", "hiruhileko"],
    "freq_semestral": ["half yearly", "semiannual", "semi annual", "seihileko"],
    "freq_anual": ["annual", "yearly", "year", "annually", "urtero", "urtekoa", "annuel"],
}

_SINONIMOS_ROLES = {
    "admin": ["admin", "administrador", "administrator", "administratzailea", "administrateur"],
    "tecnico": ["tecnico", "technician", "teknikaria", "tecnic", "technicien", "tech"],
    "almacen": ["almacen", "warehouse", "store", "storekeeper", "biltegia", "magatzem", "almacen",
                "almoxarifado", "magasin"],
}

_PALABRAS_SUELO = ["suelo", "hueco de suelo", "floor", "floor space", "ground", "zorua", "zoru hutsunea",
                   "terra", "espai a terra", "chan", "oco no chan", "chao", "espaco no chao", "sol",
                   "espace au sol"]


class ErrorLectura(Exception):
    """El fichero no se puede leer (formato no soportado, dañado...). El
    mensaje ya va traducido y listo para enseñarlo."""


# ------------------------------------------------------------------ utilidades

def normalizar(texto):
    """Minúsculas, sin tildes, sin lo que va entre paréntesis y sin signos:
    'Código (opcional)' -> 'codigo'. Para comparar cabeceras y palabras clave."""
    texto = re.sub(r"\(.*?\)", " ", str(texto or ""))
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c)).lower()
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return " ".join(texto.split())


def _clave(texto):
    """Clave para buscar nombres ya existentes sin distinguir mayúsculas."""
    return " ".join(str(texto or "").split()).casefold()


def _traducciones_de(clave):
    """Todas las traducciones de una clave en los idiomas cargados."""
    return [d[clave] for d in idiomas.TRADUCCIONES.values() if clave in d]


def etiqueta_campo(campo):
    return t(f"imp_campo_{campo}")


def texto_celda(v):
    """Valor de una celda tal y como se enseña o se guarda como texto. Un
    número entero leído del Excel como 12.0 se queda en '12' (códigos de
    artículo, logins numéricos...)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if v.time() == datetime.min.time() else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return " ".join(str(v).split())


def _numero(v):
    """Número o None si la celda está vacía. Acepta '1,5', '1.234,5' y '1,234.5'.
    Lanza ValueError si no es un número."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        raise ValueError
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    return float(s)


def _fecha(v):
    """Fecha en formato AAAA-MM-DD o None si la celda está vacía. Acepta
    celdas de fecha del Excel, AAAA-MM-DD, AAAA/MM/DD (euskera) y DD/MM/AAAA
    (resto de idiomas de la app). Lanza ValueError si no se entiende."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        # Fecha guardada como número de serie del Excel (días desde 1899-12-30)
        if 1 <= v < 200000:
            return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
        raise ValueError
    s = str(v).strip().split(" ")[0].split("T")[0]
    partes = re.split(r"[/\-.]", s)
    if len(partes) != 3 or not all(p.isdigit() for p in partes):
        raise ValueError
    if len(partes[0]) == 4:
        a, m, d = (int(p) for p in partes)
    else:
        d, m, a = (int(p) for p in partes)
        if a < 100:
            a += 2000
    return date(a, m, d).isoformat()


def _frecuencia(v):
    """Forma canónica (en español, como se guarda en la BD) de una frecuencia
    escrita en cualquier idioma, o None si no se reconoce."""
    n = normalizar(texto_celda(v))
    for clave in _CLAVES_FRECUENCIA:
        candidatos = {normalizar(x) for x in _traducciones_de(clave)} | set(_SINONIMOS_FRECUENCIA[clave])
        if n in candidatos:
            return idiomas.TRADUCCIONES["es"][clave]
    return None


def _roles(v):
    """(roles reconocidos, textos no reconocidos) a partir de una celda tipo
    'Técnico, Almacén'."""
    reconocidos, desconocidos = [], []
    for trozo in re.split(r"[,;/+|]", texto_celda(v)):
        n = normalizar(trozo)
        if not n:
            continue
        for rol, sinonimos in _SINONIMOS_ROLES.items():
            if n in {normalizar(x) for x in _traducciones_de(f"rol_{rol}")} | set(sinonimos):
                if rol not in reconocidos:
                    reconocidos.append(rol)
                break
        else:
            desconocidos.append(trozo.strip())
    return reconocidos, desconocidos


def _indice_letras(letras):
    """'A' -> 1, 'Z' -> 26, 'AA' -> 27 (inversa de almacen._letra_seccion)."""
    n = 0
    for c in letras.upper():
        n = n * 26 + (ord(c) - ord("A") + 1)
    return n


def _balda(v):
    """(número, se_dio_como_letra) de una balda: '3', 'B', 'Balda 3', 'Suelo'...
    El hueco de suelo es la balda 0. Lanza ValueError si no se entiende."""
    s = texto_celda(v)
    if not s:
        return None
    n = normalizar(s)
    suelo = set(_PALABRAS_SUELO) | {normalizar(x) for x in _traducciones_de("txt_hueco_suelo_corto")}
    if n in suelo or n == "0":
        return 0, False
    ultimo = n.split(" ")[-1] if n else ""
    if ultimo.isdigit():
        return int(ultimo), False
    if ultimo.isalpha() and len(ultimo) <= 3:
        return _indice_letras(ultimo), True
    raise ValueError


def _password_temporal():
    """Contraseña temporal fácil de dictar: sin 0/O, 1/l/I."""
    alfabeto = "abcdefghjkmnpqrstuvwxyz23456789"
    return "".join(secrets.choice(alfabeto) for _ in range(8))


# --------------------------------------------------------------------- lectura

def hojas(ruta):
    """Nombres de las hojas de un Excel ([] para un CSV)."""
    if not ruta.lower().endswith(".xlsx"):
        _comprobar_extension(ruta)
        return []
    libro = _abrir_excel(ruta)
    try:
        return list(libro.sheetnames)
    finally:
        libro.close()


def _comprobar_extension(ruta):
    r = ruta.lower()
    if r.endswith(".xls"):
        raise ErrorLectura(t("imp_err_xls"))
    if not (r.endswith(".csv") or r.endswith(".txt") or r.endswith(".xlsx")):
        raise ErrorLectura(t("imp_err_formato"))


def _abrir_excel(ruta):
    try:
        import openpyxl
        return openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except Exception as e:
        raise ErrorLectura(t("imp_err_leer").format(error=e))


def leer(ruta, hoja=None):
    """Devuelve (cabeceras, filas), con filas = [(número de fila en el fichero,
    [valores...]), ...]. La cabecera es la primera fila que no está vacía; las
    filas vacías se saltan."""
    _comprobar_extension(ruta)
    if ruta.lower().endswith(".xlsx"):
        libro = _abrir_excel(ruta)
        try:
            hoja_excel = libro[hoja] if hoja else libro.worksheets[0]
            crudas = [(i, list(fila)) for i, fila in enumerate(hoja_excel.iter_rows(values_only=True), start=1)]
        except Exception as e:
            raise ErrorLectura(t("imp_err_leer").format(error=e))
        finally:
            libro.close()
    else:
        crudas = list(enumerate(_leer_csv(ruta), start=1))

    vacia = lambda valores: all(texto_celda(v) == "" for v in valores)
    crudas = [(i, v) for i, v in crudas if not vacia(v)]
    if not crudas:
        raise ErrorLectura(t("imp_err_vacio"))
    _, cabecera = crudas[0]
    # Se recortan las columnas vacías del final (el Excel a veces arrastra
    # columnas "usadas" sin contenido).
    ancho = max(len(cabecera), *(len(v) for _, v in crudas))
    while ancho > 0 and all(len(v) < ancho or texto_celda(v[ancho - 1]) == "" for _, v in crudas):
        ancho -= 1
    cabeceras = []
    for j in range(ancho):
        valor = texto_celda(cabecera[j]) if j < len(cabecera) else ""
        cabeceras.append(valor or t("imp_columna_n").format(n=j + 1))
    filas = [(i, (list(v) + [None] * ancho)[:ancho]) for i, v in crudas[1:]]
    return cabeceras, filas


def _leer_csv(ruta):
    with open(ruta, "rb") as f:
        bruto = f.read()
    # El Excel en español guarda los CSV en Windows-1252 y con ";"; el resto,
    # en UTF-8 (a veces con BOM) y con ",".
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            texto = bruto.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ErrorLectura(t("imp_err_leer").format(error="encoding"))
    primera = next((linea for linea in texto.splitlines() if linea.strip()), "")
    separador = max((";", ",", "\t"), key=primera.count)
    return [fila for fila in csv.reader(texto.splitlines(), delimiter=separador)]


# ------------------------------------------------------------- emparejamiento

def emparejar_automatico(tipo, cabeceras):
    """{campo: índice de columna} adivinado por el nombre de las cabeceras,
    en cualquiera de los idiomas de la app. Cada columna se usa una vez."""
    normalizadas = [normalizar(c) for c in cabeceras]
    usadas, mapeo = set(), {}
    for campo, _ in CAMPOS[tipo]:
        candidatos = {normalizar(x) for x in _traducciones_de(f"imp_campo_{campo}")}
        candidatos |= {normalizar(x) for x in _SINONIMOS.get((tipo, campo), [])}
        candidatos.discard("")
        for j, cab in enumerate(normalizadas):
            if j not in usadas and cab in candidatos:
                mapeo[campo] = j
                usadas.add(j)
                break
    return mapeo


# ------------------------------------------------------------------ resultados

class ResultadoFila:
    """Resultado de una fila: estado 'crear', 'actualizar', 'saltar' o 'error',
    más los avisos (se importa, pero conviene saberlo) o el error."""

    def __init__(self, numero, valores):
        self.numero = numero
        self.valores = valores
        self.estado = None
        self.mensajes = []

    @property
    def con_avisos(self):
        return self.estado in ("crear", "actualizar") and bool(self.mensajes)


class _ErrorFila(Exception):
    pass


def resumen(resultados):
    """Cuentas para la cabecera de la vista previa y el mensaje final."""
    cuenta = {"crear": 0, "actualizar": 0, "saltar": 0, "error": 0, "avisos": 0}
    for r in resultados:
        cuenta[r.estado] += 1
        if r.con_avisos:
            cuenta["avisos"] += 1
    return cuenta


# -------------------------------------------------------------------- proceso

def procesar(tipo, filas, mapeo, modo="saltar", aplicar=False, usuario_id=None, usuario_nombre=None):
    """Importa (aplicar=True) o simula la importación (aplicar=False) de
    `filas` según `mapeo` ({campo: índice de columna}). `modo` decide qué
    hacer con lo que ya existe: 'saltar' o 'actualizar'.

    Todo va en una única transacción: o se importa todo lo válido, o nada (si
    algo falla de forma inesperada, se deshace entero y se relanza el error).
    Devuelve (resultados, extra), con extra["passwords"] = [(login, nombre,
    contraseña temporal)] de los usuarios creados."""
    ruta = almacen.RUTA_DB if tipo == "articulos" else usuarios.RUTA_DB
    con = sqlite3.connect(ruta, timeout=20, isolation_level=None)
    con.row_factory = sqlite3.Row
    if tipo == "articulos":
        con.execute("PRAGMA foreign_keys = ON")
    importador = {"maquinas": _ImportadorMaquinas, "articulos": _ImportadorArticulos,
                  "avisos": _ImportadorAvisos, "usuarios": _ImportadorUsuarios}[tipo](
        con, modo, aplicar, usuario_id, usuario_nombre)
    resultados = []
    try:
        con.execute("BEGIN IMMEDIATE")
        for numero, valores in filas:
            r = ResultadoFila(numero, valores)
            celda = lambda campo: valores[mapeo[campo]] if mapeo.get(campo) is not None else None
            if _es_fila_ejemplo(tipo, celda):
                r.estado = "saltar"
                r.mensajes = [t("imp_fila_ejemplo")]
                resultados.append(r)
                continue
            con.execute("SAVEPOINT fila")
            try:
                for campo, obligatorio in CAMPOS[tipo]:
                    if obligatorio and not texto_celda(celda(campo)):
                        raise _ErrorFila(t("imp_err_falta").format(campo=etiqueta_campo(campo)))
                importador.fila(r, celda)
                con.execute("RELEASE fila")
            except _ErrorFila as e:
                con.execute("ROLLBACK TO fila")
                con.execute("RELEASE fila")
                r.estado = "error"
                r.mensajes = [str(e)]
            resultados.append(r)
        con.execute("COMMIT" if aplicar else "ROLLBACK")
    except Exception:
        if con.in_transaction:
            con.execute("ROLLBACK")
        raise
    finally:
        con.close()
    return resultados, {"passwords": getattr(importador, "passwords", [])}


def _es_fila_ejemplo(tipo, celda):
    """True si la fila es la de ejemplo de la plantilla (en cualquier idioma,
    por si la plantilla se descargó con la app en otro): se salta sola en vez
    de importar un "Horno de fusión 1" de mentira si se olvida borrarla."""
    for idioma in idiomas.TRADUCCIONES:
        traducir = lambda clave: idiomas.TRADUCCIONES[idioma].get(clave, clave)
        textos = {c: v for c, v in _ejemplos(traducir)[tipo].items() if isinstance(v, str)}
        if all(_clave(texto_celda(celda(c))) == _clave(v) for c, v in textos.items()):
            return True
    return False


class _Importador:
    def __init__(self, con, modo, aplicar, usuario_id, usuario_nombre):
        self.con = con
        self.modo = modo
        self.aplicar = aplicar
        self.usuario_id = usuario_id
        self.usuario_nombre = usuario_nombre
        # Claves de lo tocado por filas anteriores de este mismo fichero, para
        # avisar de filas repetidas en vez de importarlas dos veces.
        self.vistos = {}

    def _repetida(self, clave, r):
        if clave in self.vistos:
            raise _ErrorFila(t("imp_err_repetida").format(fila=self.vistos[clave]))
        self.vistos[clave] = r.numero

    @staticmethod
    def _ahora():
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


class _ImportadorMaquinas(_Importador):
    def __init__(self, *args):
        super().__init__(*args)
        self.maquinas = {(f["padre_id"], _clave(f["nombre"])): f["id"]
                         for f in self.con.execute("SELECT id, nombre, padre_id FROM maquinas ORDER BY id DESC")}
        self.zonas = {_clave(f["nombre"]): f["id"] for f in self.con.execute("SELECT id, nombre FROM zonas")}
        # Máquinas superiores creadas solas porque otra fila las nombraba:
        # si más abajo aparece su propia fila, se completa (zona, notas) en vez
        # de darla por "ya existe".
        self.creadas_como_padre = set()

    def _crear(self, nombre, padre_id, notas="", zona_id=None):
        cur = self.con.execute(
            "INSERT INTO maquinas (nombre, padre_id, notas, creado, foto, zona_id) VALUES (?, ?, ?, ?, NULL, ?)",
            (nombre, padre_id, notas, datetime.now().isoformat(timespec="seconds"), zona_id))
        self.maquinas[(padre_id, _clave(nombre))] = cur.lastrowid
        return cur.lastrowid

    def fila(self, r, celda):
        tramos = [p.strip() for p in texto_celda(celda("padre")).split(">") if p.strip()]
        tramos += [p.strip() for p in texto_celda(celda("nombre")).split(">") if p.strip()]
        if not tramos:
            raise _ErrorFila(t("imp_err_falta").format(campo=etiqueta_campo("nombre")))
        notas = texto_celda(celda("notas"))
        nombre_zona = texto_celda(celda("zona"))
        avisos = []

        padre_id = None
        for tramo in tramos[:-1]:
            id_tramo = self.maquinas.get((padre_id, _clave(tramo)))
            if id_tramo is None:
                id_tramo = self._crear(tramo, padre_id)
                self.creadas_como_padre.add(id_tramo)
                avisos.append(t("imp_aviso_crear_maquina_padre").format(nombre=tramo))
            padre_id = id_tramo
        nombre = tramos[-1]

        zona_id = None
        if nombre_zona:
            zona_id = self.zonas.get(_clave(nombre_zona))
            if zona_id is None:
                zona_id = self.con.execute("INSERT INTO zonas (nombre) VALUES (?)", (nombre_zona,)).lastrowid
                self.zonas[_clave(nombre_zona)] = zona_id
                avisos.append(t("imp_aviso_crear_zona").format(nombre=nombre_zona))

        existente = self.maquinas.get((padre_id, _clave(nombre)))
        self._repetida((padre_id, _clave(nombre)), r)
        if existente is None:
            self._crear(nombre, padre_id, notas, zona_id)
            r.estado = "crear"
        elif existente in self.creadas_como_padre or self.modo == "actualizar":
            cambios, valores = [], []
            if notas:
                cambios.append("notas=?"); valores.append(notas)
            if zona_id is not None:
                cambios.append("zona_id=?"); valores.append(zona_id)
            if cambios:
                self.con.execute(f"UPDATE maquinas SET {', '.join(cambios)} WHERE id=?", (*valores, existente))
            r.estado = "crear" if existente in self.creadas_como_padre else "actualizar"
            self.creadas_como_padre.discard(existente)
        else:
            r.estado = "saltar"
            avisos = [t("imp_ya_existe")]
        r.mensajes = avisos


class _ImportadorArticulos(_Importador):
    def __init__(self, *args):
        super().__init__(*args)
        con = self.con
        self.zonas = {}
        for f in con.execute("SELECT id, nombre FROM zonas ORDER BY id"):
            self.zonas.setdefault(_clave(f["nombre"]), f["id"])
        self.estanterias = [dict(f) for f in con.execute("SELECT id, nombre, zona_id, estilo_baldas FROM estanterias")]
        self.baldas = {(f["estanteria_id"], f["numero"]): f["id"]
                       for f in con.execute("SELECT id, estanteria_id, numero FROM baldas")}
        self.materiales = [dict(f) for f in con.execute("SELECT * FROM materiales")]

    # ------------------------------------------------------------ ubicación

    def _zona(self, nombre, avisos):
        zona_id = self.zonas.get(_clave(nombre))
        if zona_id is None:
            zona_id = self.con.execute("INSERT INTO zonas (nombre) VALUES (?)", (nombre,)).lastrowid
            self.zonas[_clave(nombre)] = zona_id
            avisos.append(t("imp_aviso_crear_zona").format(nombre=nombre))
        return zona_id

    def _estanteria(self, nombre_zona, nombre_est, como_letra, avisos):
        if nombre_zona:
            zona_id = self._zona(nombre_zona, avisos)
            candidatas = [e for e in self.estanterias if e["zona_id"] == zona_id and _clave(e["nombre"]) == _clave(nombre_est)]
        else:
            candidatas = [e for e in self.estanterias if _clave(e["nombre"]) == _clave(nombre_est)]
            if len(candidatas) > 1:
                raise _ErrorFila(t("imp_err_estanteria_varias_zonas").format(nombre=nombre_est))
            if not candidatas:
                # Para crear la estantería hace falta una zona: si solo hay
                # una se usa esa; si no hay ninguna se crea una general.
                ids_zonas = sorted(set(self.zonas.values()))
                if len(ids_zonas) > 1:
                    raise _ErrorFila(t("imp_err_falta_zona").format(nombre=nombre_est))
                zona_id = ids_zonas[0] if ids_zonas else self._zona(t("imp_zona_general"), avisos)
        if candidatas:
            return candidatas[0]
        estilo = "letra" if como_letra else "numero"
        est_id = self.con.execute(
            "INSERT INTO estanterias (nombre, tiene_hueco_suelo, estilo_baldas, zona_id) VALUES (?, 0, ?, ?)",
            (nombre_est, estilo, zona_id)).lastrowid
        est = {"id": est_id, "nombre": nombre_est, "zona_id": zona_id, "estilo_baldas": estilo}
        self.estanterias.append(est)
        avisos.append(t("imp_aviso_crear_estanteria").format(nombre=nombre_est, zona=self._nombre_zona(zona_id)))
        return est

    def _nombre_zona(self, zona_id):
        fila = self.con.execute("SELECT nombre FROM zonas WHERE id=?", (zona_id,)).fetchone()
        return fila["nombre"] if fila else ""

    def _balda(self, est, numero, avisos):
        balda_id = self.baldas.get((est["id"], numero))
        if balda_id is not None:
            return balda_id
        # Se crean también las baldas de debajo que falten: una estantería con
        # la balda 3 pero sin la 1 y la 2 no tiene sentido físicamente.
        nuevas = [numero] if numero == 0 else [n for n in range(1, numero + 1) if (est["id"], n) not in self.baldas]
        for n in nuevas:
            self.baldas[(est["id"], n)] = self.con.execute(
                "INSERT INTO baldas (estanteria_id, numero) VALUES (?, ?)", (est["id"], n)).lastrowid
        if numero == 0:
            self.con.execute("UPDATE estanterias SET tiene_hueco_suelo=1 WHERE id=?", (est["id"],))
            avisos.append(t("imp_aviso_crear_suelo").format(estanteria=est["nombre"]))
        else:
            textos = ", ".join(almacen.texto_balda(n, est["estilo_baldas"]) for n in nuevas)
            avisos.append(t("imp_aviso_crear_baldas").format(baldas=textos, estanteria=est["nombre"]))
        return self.baldas[(est["id"], numero)]

    def _seccion(self, balda_id, nombre, avisos):
        secciones = [dict(f) for f in self.con.execute(
            "SELECT id, nombre, implicita FROM secciones WHERE balda_id=? ORDER BY id", (balda_id,))]
        if not nombre:
            # Sin sección: igual que almacen.seccion_para_balda (sección implícita)
            implicita = next((s for s in secciones if s["implicita"]), None)
            if implicita:
                return implicita["id"]
            return self.con.execute(
                "INSERT INTO secciones (balda_id, nombre, implicita) VALUES (?, 'Sección A', 1)", (balda_id,)).lastrowid
        buscada = _clave(nombre)
        for s in secciones:
            if _clave(s["nombre"]) == buscada or _clave(s["nombre"]).split(" ")[-1] == buscada:
                if s["implicita"]:
                    self.con.execute("UPDATE secciones SET implicita=0 WHERE id=?", (s["id"],))
                return s["id"]
        # Nombres cortos ('A', '2') siguen la misma convención que al crear las
        # secciones desde el escritorio ('Sección A'); los largos se dejan tal cual.
        nombre_final = f"Sección {nombre.upper()}" if re.fullmatch(r"[A-Za-z0-9]{1,3}", nombre) else nombre
        # Igual que almacen.crear_seccion: al crear una sección explícita, la
        # implícita de esa balda deja de estar oculta.
        self.con.execute("UPDATE secciones SET implicita=0 WHERE balda_id=? AND implicita=1", (balda_id,))
        seccion_id = self.con.execute(
            "INSERT INTO secciones (balda_id, nombre, implicita) VALUES (?, ?, 0)", (balda_id, nombre_final)).lastrowid
        avisos.append(t("imp_aviso_crear_seccion").format(nombre=nombre_final))
        return seccion_id

    def _texto_ubicacion(self, seccion_id):
        """Como almacen.obtener_ubicacion_texto, pero con esta conexión (para
        ver lo creado dentro de la transacción)."""
        if not seccion_id:
            return ""
        f = self.con.execute("""
            SELECT e.nombre AS estanteria, e.estilo_baldas AS estilo, b.numero AS balda,
                   s.nombre AS seccion, s.implicita AS implicita
            FROM secciones s JOIN baldas b ON s.balda_id = b.id JOIN estanterias e ON b.estanteria_id = e.id
            WHERE s.id = ?""", (seccion_id,)).fetchone()
        if not f:
            return ""
        balda = "Suelo" if f["balda"] == 0 else f"Balda {almacen.texto_balda(f['balda'], f['estilo'])}"
        texto = f"{f['estanteria']} > {balda}"
        return texto if f["implicita"] else f"{texto} > {f['seccion']}"

    def _movimiento(self, material_id, tipo, cantidad, motivo):
        self.con.execute("""
            INSERT INTO movimientos (material_id, fecha, tipo, cantidad, usuario_id, usuario_nombre, motivo)
            VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (material_id, self._ahora(), tipo, cantidad, self.usuario_id, self.usuario_nombre, motivo))

    # ----------------------------------------------------------------- fila

    def fila(self, r, celda):
        codigo = texto_celda(celda("codigo"))
        nombre = texto_celda(celda("nombre"))
        descripcion = texto_celda(celda("descripcion"))
        unidad = texto_celda(celda("unidad"))
        try:
            stock = _numero(celda("stock"))
        except ValueError:
            raise _ErrorFila(t("imp_err_numero").format(campo=etiqueta_campo("stock"), valor=texto_celda(celda("stock"))))
        try:
            minimo = _numero(celda("minimo"))
        except ValueError:
            raise _ErrorFila(t("imp_err_numero").format(campo=etiqueta_campo("minimo"), valor=texto_celda(celda("minimo"))))
        if (stock is not None and stock < 0) or (minimo is not None and minimo < 0):
            raise _ErrorFila(t("imp_err_negativo"))
        try:
            balda = _balda(celda("balda"))
        except ValueError:
            raise _ErrorFila(t("imp_err_balda").format(valor=texto_celda(celda("balda"))))
        nombre_zona = texto_celda(celda("zona"))
        nombre_est = texto_celda(celda("estanteria"))
        nombre_seccion = texto_celda(celda("seccion"))

        if codigo:
            existente = next((m for m in self.materiales if m["codigo"] and _clave(m["codigo"]) == _clave(codigo)), None)
            self._repetida(("codigo", _clave(codigo)), r)
        else:
            existente = next((m for m in self.materiales if _clave(m["nombre"]) == _clave(nombre)), None)
            self._repetida(("nombre", _clave(nombre)), r)

        if existente and self.modo != "actualizar":
            r.estado = "saltar"
            r.mensajes = [t("imp_ya_existe")]
            return

        avisos = []
        seccion_id = None
        ubicacion_indicada = bool(nombre_est or balda is not None or nombre_seccion)
        if nombre_est and balda is not None:
            est = self._estanteria(nombre_zona, nombre_est, balda[1], avisos)
            balda_id = self._balda(est, balda[0], avisos)
            seccion_id = self._seccion(balda_id, nombre_seccion, avisos)
        elif ubicacion_indicada:
            avisos.append(t("imp_aviso_ubicacion_incompleta"))

        if existente is None:
            cur = self.con.execute("""
                INSERT INTO materiales (codigo, nombre, descripcion, unidad, stock_actual, stock_minimo, seccion_id, foto)
                VALUES (?, ?, ?, ?, ?, ?, ?, NULL)""",
                (codigo, nombre, descripcion, unidad, stock or 0, minimo or 0, seccion_id))
            if stock:
                self._movimiento(cur.lastrowid, "entrada", stock, almacen.MOTIVO_ALTA_INICIAL)
            self.materiales.append({"id": cur.lastrowid, "codigo": codigo, "nombre": nombre})
            r.estado = "crear"
        else:
            # Actualizar: solo los campos que vienen con valor; lo vacío no borra nada.
            nuevos = {"codigo": codigo, "nombre": nombre, "descripcion": descripcion, "unidad": unidad}
            cambios = {k: v for k, v in nuevos.items() if v and v != (existente.get(k) or "")}
            if minimo is not None:
                cambios["stock_minimo"] = minimo
            seccion_anterior = existente.get("seccion_id")
            if seccion_id and seccion_id != seccion_anterior:
                origen = self._texto_ubicacion(seccion_anterior) or "Sin ubicación"
                cambios["seccion_id"] = seccion_id
            if cambios:
                asignaciones = ", ".join(f"{k}=?" for k in cambios)
                self.con.execute(f"UPDATE materiales SET {asignaciones} WHERE id=?", (*cambios.values(), existente["id"]))
            if "seccion_id" in cambios:
                self.con.execute("""
                    INSERT INTO movimientos (material_id, fecha, tipo, cantidad, usuario_id, usuario_nombre, motivo)
                    VALUES (?, ?, 'traslado', 0, ?, ?, ?)""",
                    (existente["id"], self._ahora(), self.usuario_id, self.usuario_nombre,
                     f"De {origen} a {self._texto_ubicacion(seccion_id)}"))
                # Igual que almacen._borrar_seccion_si_implicita_y_vacia
                if seccion_anterior:
                    self.con.execute("""
                        DELETE FROM secciones WHERE id=? AND implicita=1
                        AND NOT EXISTS (SELECT 1 FROM materiales WHERE seccion_id=?)""",
                        (seccion_anterior, seccion_anterior))
            if stock is not None and stock != existente["stock_actual"]:
                # El stock nunca se pisa sin más: la diferencia queda como
                # movimiento para que el historial cuadre.
                diferencia = stock - existente["stock_actual"]
                self.con.execute("UPDATE materiales SET stock_actual=? WHERE id=?", (stock, existente["id"]))
                self._movimiento(existente["id"], "entrada" if diferencia > 0 else "salida",
                                 abs(diferencia), MOTIVO_AJUSTE_IMPORTACION)
                avisos.append(t("imp_aviso_ajuste_stock").format(antes=f"{existente['stock_actual']:g}", despues=f"{stock:g}"))
            r.estado = "actualizar"
        r.mensajes = avisos


class _ImportadorAvisos(_Importador):
    def __init__(self, *args):
        super().__init__(*args)
        self.avisos = {_clave(f["titulo"]): f["id"] for f in self.con.execute("SELECT id, titulo FROM avisos_recurrentes")}
        self.usuarios = [dict(f) for f in self.con.execute("SELECT id, login, nombre FROM usuarios WHERE activo=1")]
        self.especialidades = {_clave(f["nombre"]): f["id"] for f in self.con.execute("SELECT id, nombre FROM especialidades")}

    def _especialidad(self, nombre, avisos):
        esp_id = self.especialidades.get(_clave(nombre))
        if esp_id is None:
            esp_id = self.con.execute("INSERT INTO especialidades (nombre) VALUES (?)", (nombre,)).lastrowid
            self.especialidades[_clave(nombre)] = esp_id
            avisos.append(t("imp_aviso_crear_especialidad").format(nombre=nombre))
        return esp_id

    def fila(self, r, celda):
        titulo = texto_celda(celda("titulo"))
        try:
            inicio = _fecha(celda("fecha_inicio"))
        except ValueError:
            raise _ErrorFila(t("imp_err_fecha").format(campo=etiqueta_campo("fecha_inicio"), valor=texto_celda(celda("fecha_inicio"))))
        try:
            ultima = _fecha(celda("ultima")) or ""
        except ValueError:
            raise _ErrorFila(t("imp_err_fecha").format(campo=etiqueta_campo("ultima"), valor=texto_celda(celda("ultima"))))
        frecuencia = _frecuencia(celda("frecuencia"))
        if not frecuencia:
            raise _ErrorFila(t("imp_err_frecuencia").format(valor=texto_celda(celda("frecuencia"))))
        try:
            margen = _numero(celda("margen"))
        except ValueError:
            margen = -1
        if margen is not None and (margen != int(margen) or not 1 <= margen <= 365):
            raise _ErrorFila(t("imp_err_margen").format(valor=texto_celda(celda("margen"))))

        avisos = []
        asignado_a = asignado_nombre = None
        tecnico = texto_celda(celda("tecnico"))
        if tecnico:
            u = next((u for u in self.usuarios if _clave(u["login"]) == _clave(tecnico)), None) or \
                next((u for u in self.usuarios if _clave(u["nombre"]) == _clave(tecnico)), None)
            if u:
                asignado_a, asignado_nombre = u["id"], u["nombre"]
            else:
                avisos.append(t("imp_aviso_tecnico_no_existe").format(nombre=tecnico))
        nombre_esp = texto_celda(celda("especialidad"))

        self._repetida(_clave(titulo), r)
        existente = self.avisos.get(_clave(titulo))
        if existente and self.modo != "actualizar":
            r.estado = "saltar"
            r.mensajes = [t("imp_ya_existe")]
            return
        especialidad_id = self._especialidad(nombre_esp, avisos) if nombre_esp else None

        if existente is None:
            self.con.execute("""
                INSERT INTO avisos_recurrentes (titulo, fecha_inicio, frecuencia, duracion_dias, ultima_completada,
                                                asignado_a, asignado_nombre, especialidad_id)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                (titulo, inicio, frecuencia, int(margen or MARGEN_POR_DEFECTO), ultima,
                 asignado_a, asignado_nombre, especialidad_id))
            r.estado = "crear"
        else:
            cambios = {"fecha_inicio": inicio, "frecuencia": frecuencia}
            if margen is not None:
                cambios["duracion_dias"] = int(margen)
            if ultima:
                cambios["ultima_completada"] = ultima
            if asignado_a is not None:
                cambios["asignado_a"], cambios["asignado_nombre"] = asignado_a, asignado_nombre
            if especialidad_id is not None:
                cambios["especialidad_id"] = especialidad_id
            asignaciones = ", ".join(f"{k}=?" for k in cambios)
            self.con.execute(f"UPDATE avisos_recurrentes SET {asignaciones} WHERE id=?", (*cambios.values(), existente))
            r.estado = "actualizar"
        r.mensajes = avisos


class _ImportadorUsuarios(_ImportadorAvisos):
    def __init__(self, *args):
        super().__init__(*args)
        self.todos = {_clave(f["login"]): dict(f) for f in self.con.execute("SELECT id, login, nombre, activo FROM usuarios")}
        self.passwords = []

    def _admins_activos(self):
        return {f["usuario_id"] for f in self.con.execute("""
            SELECT ur.usuario_id FROM usuario_roles ur JOIN usuarios u ON u.id = ur.usuario_id
            WHERE ur.rol='admin' AND u.activo=1""")}

    def fila(self, r, celda):
        login = texto_celda(celda("login"))
        nombre = texto_celda(celda("nombre"))
        roles, desconocidos = _roles(celda("roles"))
        avisos = [t("imp_aviso_rol_desconocido").format(rol=d) for d in desconocidos]
        nombres_esp = [e.strip() for e in re.split(r"[,;/|]", texto_celda(celda("especialidad"))) if e.strip()]

        self._repetida(_clave(login), r)
        existente = self.todos.get(_clave(login))
        if existente and self.modo != "actualizar":
            r.estado = "saltar"
            r.mensajes = [t("imp_ya_existe")]
            return
        ids_esp = [self._especialidad(n, avisos) for n in nombres_esp]

        if existente is None:
            roles = roles or ["tecnico"]
            password = _password_temporal()
            # El hash (PBKDF2, lento a propósito) solo hace falta al importar de
            # verdad; en la vista previa se usa un marcador.
            hash_pw = usuarios.hash_password(password) if self.aplicar else "-"
            cur = self.con.execute("""
                INSERT INTO usuarios (login, nombre, password_hash, rol, activo, debe_cambiar, creado)
                VALUES (?, ?, ?, ?, 1, 1, ?)""",
                (login, nombre, hash_pw, "admin" if "admin" in roles else roles[0],
                 datetime.now().isoformat(timespec="seconds")))
            usuario_id = cur.lastrowid
            r.estado = "crear"
            self.passwords.append((login, nombre, password))
        else:
            usuario_id = existente["id"]
            self.con.execute("UPDATE usuarios SET nombre=? WHERE id=?", (nombre, usuario_id))
            if roles and "admin" not in roles and self._admins_activos() == {usuario_id}:
                raise _ErrorFila(t("imp_err_ultimo_admin"))
            if roles:
                self.con.execute("DELETE FROM usuario_roles WHERE usuario_id=?", (usuario_id,))
                self.con.execute("UPDATE usuarios SET rol=? WHERE id=?",
                                 ("admin" if "admin" in roles else roles[0], usuario_id))
            r.estado = "actualizar"
        if roles:
            self.con.executemany("INSERT OR IGNORE INTO usuario_roles (usuario_id, rol) VALUES (?, ?)",
                                 [(usuario_id, rol) for rol in roles])
        if ids_esp:
            usuarios.establecer_especialidades_usuario(usuario_id, ids_esp, con=self.con)
        r.mensajes = avisos


# ------------------------------------------------------------ ficheros Excel

def _ejemplos(t=t):
    """Fila de ejemplo de cada hoja de la plantilla, en el idioma actual (o
    con la función de traducción que se pase)."""
    return {
        "maquinas": {"nombre": t("imp_ej_maquina"), "padre": t("imp_ej_maquina_padre"),
                     "zona": t("imp_ej_zona_maquina"), "notas": t("imp_ej_notas")},
        "articulos": {"codigo": "TOR-M8", "nombre": t("imp_ej_articulo"), "descripcion": t("imp_ej_descripcion"),
                      "unidad": t("imp_ej_unidad"), "stock": 120, "minimo": 50,
                      "zona": t("imp_ej_zona_almacen"), "estanteria": t("imp_ej_estanteria"),
                      "balda": 2, "seccion": "A"},
        "avisos": {"titulo": t("imp_ej_aviso"), "fecha_inicio": date(date.today().year, 1, 15),
                   "frecuencia": t("freq_mensual"), "margen": MARGEN_POR_DEFECTO,
                   "tecnico": "jperez", "especialidad": t("imp_ej_especialidad"), "ultima": None},
        "usuarios": {"login": "jperez", "nombre": "Juan Pérez", "roles": t("rol_tecnico"),
                     "especialidad": t("imp_ej_especialidad")},
    }


def nombre_hoja(tipo):
    """Nombre de la hoja de un tipo en la plantilla (los nombres de hoja del
    Excel tienen como mucho 31 caracteres y no admiten ciertos signos)."""
    return re.sub(r"[\[\]:*?/\\]", "", t(f"imp_tipo_{tipo}"))[:31]


def generar_plantilla(ruta):
    """Excel con una hoja de instrucciones y una hoja por tipo (cabeceras en el
    idioma actual, fila de ejemplo y desplegables en las columnas cerradas)."""
    import xlsxwriter
    libro = xlsxwriter.Workbook(ruta)
    negrita = libro.add_format({"bold": True, "font_size": 14})
    texto = libro.add_format({"text_wrap": True, "valign": "top"})
    cab = libro.add_format({"bold": True, "bg_color": "#3daee9", "font_color": "white", "border": 1})
    cab_obligatoria = libro.add_format({"bold": True, "bg_color": "#c0392b", "font_color": "white", "border": 1})
    ejemplo = libro.add_format({"italic": True, "font_color": "#555555"})
    fecha = libro.add_format({"italic": True, "font_color": "#555555", "num_format": "yyyy-mm-dd"})

    hoja = libro.add_worksheet(re.sub(r"[\[\]:*?/\\]", "", t("imp_hoja_instrucciones"))[:31])
    hoja.set_column(0, 0, 110)
    hoja.write(0, 0, t("imp_instr_titulo"), negrita)
    for i in range(1, 8):
        hoja.write(i + 1, 0, t(f"imp_instr_{i}"), texto)

    ejemplos = _ejemplos()
    for tipo in TIPOS:
        h = libro.add_worksheet(nombre_hoja(tipo))
        h.freeze_panes(1, 0)
        for j, (campo, obligatorio) in enumerate(CAMPOS[tipo]):
            h.set_column(j, j, 22)
            h.write(0, j, etiqueta_campo(campo), cab_obligatoria if obligatorio else cab)
            valor = ejemplos[tipo].get(campo)
            if isinstance(valor, date):
                h.write_datetime(1, j, datetime.combine(valor, datetime.min.time()), fecha)
            elif valor is not None:
                h.write(1, j, valor, ejemplo)
            if campo == "frecuencia":
                h.data_validation(1, j, 2000, j, {"validate": "list", "source": [t(c) for c in reversed(_CLAVES_FRECUENCIA)]})
            elif campo in ("fecha_inicio", "ultima"):
                h.set_column(j, j, 22, libro.add_format({"num_format": "yyyy-mm-dd"}))
            elif campo == "roles":
                h.data_validation(1, j, 2000, j, {"validate": "list", "error_type": "information",
                                                  "source": [t("rol_tecnico"), t("rol_almacen"), t("rol_admin")]})
    libro.close()


def exportar_errores(ruta, cabeceras, resultados):
    """Excel con las filas que no se han podido importar, sus columnas
    originales y el motivo, para corregirlas y volver a importarlas."""
    import xlsxwriter
    libro = xlsxwriter.Workbook(ruta)
    h = libro.add_worksheet()
    cab = libro.add_format({"bold": True, "bg_color": "#3daee9", "font_color": "white"})
    rojo = libro.add_format({"font_color": "#c0392b"})
    for j, c in enumerate(list(cabeceras) + [t("imp_col_error")]):
        h.write(0, j, c, cab)
        h.set_column(j, j, 20)
    i = 1
    for r in resultados:
        if r.estado != "error":
            continue
        for j, v in enumerate(r.valores):
            if isinstance(v, (datetime, date)):
                h.write(i, j, texto_celda(v))
            elif v is not None:
                h.write(i, j, v)
        h.write(i, len(cabeceras), " · ".join(r.mensajes), rojo)
        i += 1
    h.set_column(len(cabeceras), len(cabeceras), 60)
    libro.close()


def exportar_passwords(ruta, passwords):
    """Excel con el usuario, el nombre y la contraseña temporal de cada
    usuario creado (tendrá que cambiarla al entrar por primera vez)."""
    import xlsxwriter
    libro = xlsxwriter.Workbook(ruta)
    h = libro.add_worksheet()
    cab = libro.add_format({"bold": True, "bg_color": "#3daee9", "font_color": "white"})
    for j, c in enumerate((etiqueta_campo("login"), etiqueta_campo("nombre"), t("imp_col_password"))):
        h.write(0, j, c, cab)
        h.set_column(j, j, 26)
    for i, fila in enumerate(passwords, start=1):
        for j, v in enumerate(fila):
            h.write_string(i, j, v)
    libro.close()
